import fitz
import re
import os
import logging
import tkinter as tk
from tkinter import filedialog, messagebox, ttk
from typing import Dict, Set

try:
    import docx
except ImportError:
    docx = None

try:
    import pyi_splash
    pyi_splash.close()
except ImportError:
    pass    

# Configuração de logger em nível de módulo[cite: 1]
logger = logging.getLogger(__name__)
logger.setLevel(logging.INFO)
if not logger.handlers:
    handler = logging.StreamHandler()
    formatter = logging.Formatter('%(asctime)s - %(levelname)s - %(message)s')
    handler.setFormatter(formatter)
    logger.addHandler(handler)

# Configuração do NLP (spaCy)[cite: 1]
import spacy
from spacy.util import is_package
nlp = None
try:
    nlp = spacy.load("pt_core_news_lg", disable=["lemmatizer", "textcat", "custom"])
    logger.info("Modelo spaCy (pt_core_news_lg) carregado com sucesso.")
except Exception as e:
    logger.error(f"Falha ao carregar modelo spaCy: {e}")


class DocumentAnonimizer:
    # Expressões Regulares de Documentos[cite: 1]
    REGEX_CPF = re.compile(r'\b\d{3}\.\d{3}\.\d{3}-\d{2}\b')
    REGEX_RG_CONTEXTO = re.compile(r'(RG:)\s*([0-9.-]+\s?\d*)', re.IGNORECASE)
    REGEX_DATA_EXP = re.compile(r'(Data de Expedição:)\s*(\d{2}/\d{2}/\d{4})', re.IGNORECASE)
    REGEX_7_DIGITOS = re.compile(r'\b\d{7}\b')
    REGEX_CRC = re.compile(r'\b[0-9A-F]{8}\b', re.IGNORECASE)
    REGEX_EMAIL = re.compile(r'\b[A-Za-z0-9._%+-]+@[A-Za-z0-9.-]+\.[A-Z|a-z]{2,}\b')
    REGEX_TELEFONE = re.compile(r'(?:\(?\d{2}\)?\s?)?(?:9\d{4}|\d{4})[-\s]\d{4}')
    
    # Regex de resgate para forçar o spaCy a avaliar blocos suspeitos isoladamente[cite: 1]
    REGEX_NOME_CAPS = re.compile(r'\b[A-ZÀ-Ÿ]{2,}(?: [A-ZÀ-Ÿ]{2,}| DE| DA| DO| DOS| E)+\b')
    REGEX_NOME_TITLE = re.compile(r'\b[A-ZÀ-Ÿ][a-zà-ÿ]+(?: [A-ZÀ-Ÿ][a-zà-ÿ]+| de| da| do| dos| e)+\b')

    def __init__(self, opcoes: Dict[str, bool] = None):
        self.opcoes = opcoes or {
            "cpf": True,
            "rg": True,
            "email_tel": True,
            "doc_sei": True,
            "qr_code": True,
            "links": True,
            "nomes": False,
            "modo_nomes": "iniciais" 
        }

    def __enter__(self):
        return self

    def __exit__(self, exc_type, exc_val, exc_tb):
        pass

    @staticmethod
    def mascarar_cpf_fmt(cpf_original: str) -> str:
        numeros = re.sub(r'\D', '', cpf_original)
        if len(numeros) == 11:
            return f"***.{numeros[3:6]}.{numeros[6:9]}-**"
        return "***.***.***-**"

    @staticmethod
    def mascarar_doc_7_digitos_fmt(numero: str) -> str:
        if len(numero) == 7:
            return "XXXXX" + numero[-2:]
        return "XXXXXXX"

    def _gerar_iniciais(self, nome_completo: str) -> str:
        palavras = nome_completo.split()
        iniciais = []
        stopwords = {"DE", "DA", "DO", "DAS", "DOS", "E"}
        
        for p in palavras:
            p_upper = p.upper()
            if p_upper in stopwords:
                continue
            
            if len(p) <= 2 and (p.endswith('.') or p.isupper()):
                iniciais.append(p if p.endswith('.') else p + ".")
            elif len(p) > 1:
                iniciais.append(p[0].upper() + ".")
                
        return " ".join(iniciais) if iniciais else "[N.I.]"

    def extrair_nomes_nlp(self, texto: str) -> Set[str]:
        if not nlp: return set()
        
        nomes_encontrados = set()
        termos_protegidos = {"DCTI", "D.C.", "D. C.", "CNPQ", "SEI", "COORDENAÇÃO", "DIRETORIA"}
        
        doc_spacy = nlp(texto)
        for ent in doc_spacy.ents:
            if ent.label_ == "PER":
                nome_limpo = " ".join(ent.text.replace("\n", " ").split())
                if nome_limpo.upper() not in termos_protegidos:
                    if len(nome_limpo.split()) >= 2 and len(nome_limpo) > 4:
                        nomes_encontrados.add(nome_limpo)

        for padrao in [self.REGEX_NOME_CAPS, self.REGEX_NOME_TITLE]:
            for match in padrao.finditer(texto):
                snippet_original = match.group()
                nome_candidato = " ".join(snippet_original.replace("\n", " ").split())
                
                if nome_candidato in nomes_encontrados or nome_candidato.upper() in termos_protegidos:
                    continue
                    
                doc_snippet = nlp(snippet_original.title())
                for ent in doc_snippet.ents:
                    if ent.label_ == "PER":
                        if len(nome_candidato.split()) >= 2 and len(nome_candidato) > 4:
                            nomes_encontrados.add(nome_candidato)
                        
        return nomes_encontrados

    def processar_arquivo(self, caminho_entrada: str, caminho_saida: str = None) -> str:
        extensao = caminho_entrada.lower().split('.')[-1]
        
        if not caminho_saida:
            pasta_pai = os.path.dirname(caminho_entrada)
            nome_arquivo = os.path.basename(caminho_entrada)
            caminho_saida = os.path.join(pasta_pai, os.path.splitext(nome_arquivo)[0] + f"_anonimizado.{extensao}")

        logger.info(f"Iniciando anonimização: {os.path.basename(caminho_entrada)}")

        if extensao == 'pdf':
            return self._processar_pdf(caminho_entrada, caminho_saida)
        elif extensao == 'txt':
            return self._processar_txt(caminho_entrada, caminho_saida)
        elif extensao == 'docx':
            if not docx:
                raise ImportError("A biblioteca python-docx não está instalada. Execute: pip install python-docx")
            return self._processar_docx(caminho_entrada, caminho_saida)
        else:
            raise ValueError(f"Formato de arquivo não suportado: {extensao}")

    def _aplicar_substituicoes_texto(self, texto: str) -> str:
        if not texto or not texto.strip():
            return texto

        nomes_na_pagina = set()
        if self.opcoes.get("nomes") and nlp:
            nomes_na_pagina = self.extrair_nomes_nlp(texto)

        if self.opcoes.get("email_tel"):
            texto = self.REGEX_EMAIL.sub("[E-MAIL]", texto)
            
            def repl_tel(m):
                fone = m.group()
                if re.match(r'^(19|20)\d{2}-\d{4}', fone): return fone
                return "[TEL]"
            texto = self.REGEX_TELEFONE.sub(repl_tel, texto)

        if self.opcoes.get("cpf"):
            def repl_cpf(m):
                return self.mascarar_cpf_fmt(m.group())
            texto = self.REGEX_CPF.sub(repl_cpf, texto)

        if self.opcoes.get("rg"):
            texto = self.REGEX_RG_CONTEXTO.sub(r'\1 XXXXXX', texto)
            texto = self.REGEX_DATA_EXP.sub(r'\1 XX/XX/XXXX', texto)

        if self.opcoes.get("doc_sei"):
            def repl_sei(m):
                return self.mascarar_doc_7_digitos_fmt(m.group())
            texto = self.REGEX_7_DIGITOS.sub(repl_sei, texto)
            texto = self.REGEX_CRC.sub("XXXXXXXX", texto)

        if self.opcoes.get("nomes") and nlp:
            for nome in sorted(nomes_na_pagina, key=len, reverse=True):
                substituto = self._gerar_iniciais(nome) if self.opcoes.get("modo_nomes") == "iniciais" else "[NOME]"
                texto = texto.replace(nome, substituto)

        return texto

    def _processar_txt(self, caminho_entrada: str, caminho_saida: str) -> str:
        try:
            with open(caminho_entrada, 'r', encoding='utf-8') as f:
                texto = f.read()
            
            texto_anonimizado = self._aplicar_substituicoes_texto(texto)
            
            with open(caminho_saida, 'w', encoding='utf-8') as f:
                f.write(texto_anonimizado)
            return caminho_saida
        except Exception as e:
            logger.error(f"Falha de processamento TXT em {caminho_entrada}: {e}", exc_info=True)
            raise

    def _processar_docx(self, caminho_entrada: str, caminho_saida: str) -> str:
        try:
            documento = docx.Document(caminho_entrada)
            
            for p in documento.paragraphs:
                if p.text:
                    p.text = self._aplicar_substituicoes_texto(p.text)
            
            for table in documento.tables:
                for row in table.rows:
                    for cell in row.cells:
                        for p in cell.paragraphs:
                            if p.text:
                                p.text = self._aplicar_substituicoes_texto(p.text)
                                
            documento.save(caminho_saida)
            return caminho_saida
        except Exception as e:
            logger.error(f"Falha de processamento DOCX em {caminho_entrada}: {e}", exc_info=True)
            raise

    def _processar_pdf(self, caminho_entrada: str, caminho_saida: str) -> str:
        doc = None
        try:
            doc = fitz.open(caminho_entrada)

            for page in doc:
                texto_pagina = page.get_text()
                altura_pagina = page.rect.height
                largura_pagina = page.rect.width

                if self.opcoes.get("links"):
                    if hasattr(page, "delete_links"):
                        page.delete_links()
                    else:
                        for link in page.get_links():
                            page.delete_link(link)

                if self.opcoes.get("qr_code"):
                    rect_barra_lateral = fitz.Rect(largura_pagina - 25, 0, largura_pagina, altura_pagina)
                    annot = page.add_redact_annot(rect_barra_lateral, fill=(0, 0, 0))
                    annot.update()

                    areas_auth = page.search_for("A autenticidade do documento pode ser conferida no site")
                    if areas_auth:
                        for area in areas_auth:
                            rect_qr_left = fitz.Rect(area.x0 - 60, area.y0 - 10, area.x0, area.y1 + 40)
                            annot = page.add_redact_annot(rect_qr_left, fill=(1, 1, 1))
                            annot.update()

                if self.opcoes.get("email_tel"):
                    for match in self.REGEX_EMAIL.finditer(texto_pagina):
                        areas = page.search_for(match.group())
                        for area in areas:
                            tamanho = max(round(area.height * 0.75, 1), 6)
                            annot = page.add_redact_annot(area, text="[E-MAIL]", fontname="Helvetica", fontsize=tamanho, align=0, fill=(1,1,1))
                            annot.update()

                    for match in self.REGEX_TELEFONE.finditer(texto_pagina):
                        fone = match.group()
                        if re.match(r'^(19|20)\d{2}-\d{4}', fone): continue
                        
                        areas = page.search_for(fone)
                        for area in areas:
                            tamanho = max(round(area.height * 0.75, 1), 6)
                            annot = page.add_redact_annot(area, text="[TEL]", fontname="Helvetica", fontsize=tamanho, align=0, fill=(1,1,1))
                            annot.update()

                if self.opcoes.get("cpf"):
                    for match in self.REGEX_CPF.finditer(texto_pagina):
                        cpf = match.group()
                        areas = page.search_for(cpf)
                        mascara = self.mascarar_cpf_fmt(cpf)
                        for area in areas:
                            tamanho = max(round(area.height * 0.75, 1), 6)
                            annot = page.add_redact_annot(area, text=mascara, fontname="Helvetica", fontsize=tamanho, align=0, fill=(1,1,1))
                            annot.update()

                if self.opcoes.get("rg"):
                    for match in self.REGEX_RG_CONTEXTO.finditer(texto_pagina):
                        areas = page.search_for(match.group(0))
                        for area in areas:
                            tamanho = max(round(area.height * 0.75, 1), 6)
                            annot = page.add_redact_annot(area, text="RG: XXXXXX", fontname="Helvetica", fontsize=tamanho, align=0, fill=(1,1,1))
                            annot.update()
                    
                    for match in self.REGEX_DATA_EXP.finditer(texto_pagina):
                        areas = page.search_for(match.group(0))
                        for area in areas:
                            tamanho = max(round(area.height * 0.75, 1), 6)
                            annot = page.add_redact_annot(area, text="Data de Expedição: XX/XX/XXXX", fontname="Helvetica", fontsize=tamanho, align=0, fill=(1,1,1))
                            annot.update()

                if self.opcoes.get("doc_sei"):
                    limite_y_corte = altura_pagina * 0.85
                    areas_assinatura = page.search_for("Documento assinado eletronicamente por")
                    if areas_assinatura:
                        limite_y_corte = min([r.y0 for r in areas_assinatura])

                    for match in self.REGEX_7_DIGITOS.finditer(texto_pagina):
                        numero = match.group()
                        areas_num = page.search_for(numero)
                        for area in areas_num:
                            if area.y0 > limite_y_corte:
                                tamanho = max(round(area.height * 0.75, 1), 6)
                                mascara_parcial = self.mascarar_doc_7_digitos_fmt(numero)
                                annot = page.add_redact_annot(area, text=mascara_parcial, fontname="Helvetica", fontsize=tamanho, align=1, fill=(1,1,1))
                                annot.update()
                    
                    areas_auth = page.search_for("A autenticidade do documento pode ser conferida no site")
                    if areas_auth:
                        y_auth = min([r.y0 for r in areas_auth])
                        for match_crc in self.REGEX_CRC.finditer(texto_pagina):
                            crc = match_crc.group()
                            areas_crc = page.search_for(crc)
                            for area in areas_crc:
                                if area.y0 >= y_auth:
                                    tamanho = max(round(area.height * 0.75, 1), 6)
                                    annot = page.add_redact_annot(area, text="XXXXXXXX", fontname="Helvetica", fontsize=tamanho, align=1, fill=(0,0,0))
                                    annot.update()

                if self.opcoes.get("nomes") and nlp:
                    nomes_na_pagina = self.extrair_nomes_nlp(texto_pagina)
                    
                    for nome_original in nomes_na_pagina:
                        areas_nome = page.search_for(nome_original)
                        if not areas_nome:
                            continue
                            
                        areas_consolidadas = []
                        for area in areas_nome:
                            sobreposta = False
                            for i, area_cons in enumerate(areas_consolidadas):
                                if abs(area.y0 - area_cons.y0) < 2 and abs(area.y1 - area_cons.y1) < 2:
                                    novo_x0 = min(area.x0, area_cons.x0)
                                    novo_x1 = max(area.x1, area_cons.x1)
                                    areas_consolidadas[i] = fitz.Rect(novo_x0, area.y0, novo_x1, area.y1)
                                    sobreposta = True
                                    break
                            
                            if not sobreposta:
                                areas_consolidadas.append(area)

                        if self.opcoes.get("modo_nomes") == "iniciais":
                            texto_substituto = self._gerar_iniciais(nome_original)
                        else:
                            texto_substituto = "[NOME]"
                            
                        for area in areas_consolidadas:
                            tamanho = max(round(area.height * 0.75, 1), 6)
                            annot = page.add_redact_annot(
                                area, 
                                text=texto_substituto, 
                                fontname="Helvetica", 
                                fontsize=tamanho, 
                                align=0, 
                                fill=(1,1,1)
                            )
                            annot.update()

                page.apply_redactions()

            doc.save(caminho_saida, garbage=4, deflate=True)
            return caminho_saida

        except Exception as e:
            logger.error(f"Falha de processamento em {caminho_entrada}: {e}", exc_info=True)
            raise
        finally:
            if doc:
                doc.close()


class DialogoOpcoes(tk.Toplevel):
    def __init__(self, parent):
        super().__init__(parent)
        self.title("Configurações de Anonimização")
        self.geometry("750x720")
        self.resultado = None

        self.vars = {
            "cpf": tk.BooleanVar(value=True),
            "rg": tk.BooleanVar(value=True),
            "email_tel": tk.BooleanVar(value=True),
            "doc_sei": tk.BooleanVar(value=True),
            "qr_code": tk.BooleanVar(value=True),
            "links": tk.BooleanVar(value=True),
            "nomes": tk.BooleanVar(value=False),
            "modo_nomes": tk.StringVar(value="iniciais") 
        }

        self.btn_frame = ttk.Frame(self)
        self.btn_frame.pack(side=tk.BOTTOM, fill=tk.X, pady=15, padx=20)
        
        self.btn_cancel = ttk.Button(self.btn_frame, text="Cancelar", command=self.cancelar)
        self.btn_cancel.pack(side=tk.LEFT)

        self.btn_dir = ttk.Button(self.btn_frame, text="Tratar uma Pasta", command=lambda: self.confirmar("diretorio"))
        self.btn_dir.pack(side=tk.RIGHT, padx=5)

        self.btn_arq = ttk.Button(self.btn_frame, text="Tratar um Arquivo", command=lambda: self.confirmar("arquivo"))
        self.btn_arq.pack(side=tk.RIGHT, padx=5)

        self.main_frame = ttk.Frame(self)
        self.main_frame.pack(side=tk.TOP, fill=tk.BOTH, expand=True)

        self.canvas = tk.Canvas(self.main_frame)
        self.scrollbar = ttk.Scrollbar(self.main_frame, orient="vertical", command=self.canvas.yview)
        self.scrollable_frame = ttk.Frame(self.canvas)

        self.scrollable_frame.bind(
            "<Configure>",
            lambda e: self.canvas.configure(scrollregion=self.canvas.bbox("all"))
        )

        self.window_id = self.canvas.create_window((0, 0), window=self.scrollable_frame, anchor="nw")
        
        self.canvas.bind(
            "<Configure>",
            lambda e: self.canvas.itemconfig(self.window_id, width=e.width)
        )
        
        self.canvas.configure(yscrollcommand=self.scrollbar.set)
        self.canvas.pack(side=tk.LEFT, fill=tk.BOTH, expand=True)
        self.scrollbar.pack(side=tk.RIGHT, fill=tk.Y)

        frame_instrucoes = ttk.LabelFrame(self.scrollable_frame, text="Instruções e Avisos de Segurança", padding="10")
        frame_instrucoes.pack(fill=tk.X, padx=10, pady=10)

        texto_instrucoes = (
            "Anonimizer - Anonimizador de Documentos para Órgãos Públicos, versão 1.0.0.\n\n"
            "Programa desenvolvido por Alberto de Campos e Silva.\n"
            "Este software é de distribuição livre, licenciado sob a Licença Apache 2.0.\n\n"
            "Aviso de Segurança:\n"
            "1. A execução de programas não homologados pela equipe de Tecnologia da Informação representa risco à segurança digital. "
            "Ao executar este software, o usuário assume a confiança na integridade do código original desenvolvido pelo autor.\n"
            "2. Não execute este programa caso tenha sido recebido por meio de terceiros. O código-fonte oficial e o executor confiável "
            "devem ser obtidos exclusivamente no repositório: https://github.com/AlbertoCamposSilva/sei-anonimizer\n\n"
            "Isenção de Responsabilidade:\n"
            "O autor não assume responsabilidade por danos decorrentes de modificações maliciosas realizadas por terceiros, erros de operação, "
            "vulnerabilidades do sistema operacional ou qualquer incidente cibernético associado à execução deste software. "
            "Mantenha o sistema operacional e o antivírus atualizados.\n\n"
            "Instruções de Uso:\n"
            "1. Selecione as opções e os elementos que deseja mascarar ou remover dos documentos.\n"
            "2. Utilize os botões no rodapé para processar um único arquivo ou uma pasta inteira (em lote).\n"
            "3. O programa salvará automaticamente as cópias tratadas na mesma pasta do arquivo original, adicionando o sufixo '_anonimizado' ao nome do arquivo."
)
        
        lbl_instrucoes = ttk.Label(frame_instrucoes, text=texto_instrucoes, justify=tk.LEFT)
        lbl_instrucoes.pack(fill=tk.BOTH, expand=True)
        frame_instrucoes.bind("<Configure>", lambda e: lbl_instrucoes.configure(wraplength=e.width - 20))

        lbl = ttk.Label(self.scrollable_frame, text="Selecione os elementos para anonimizar:", font=("Arial", 10, "bold"))
        lbl.pack(anchor="w", padx=20, pady=10)

        opcoes_ui = [
            ("Mascarar CPFs (***.123.456-**)", "cpf"),
            ("Mascarar RGs e Datas (RG: XXXXXX)", "rg"),
            ("Mascarar E-mails e Telefones", "email_tel"),
            ("Mascarar Nº Documento/Verificador (XXXXX11)", "doc_sei"),
            ("Mascarar QR Codes e Barras Laterais (SEI)", "qr_code"),
            ("Remover todos os Links Clicáveis", "links")
        ]

        for texto, chave in opcoes_ui:
            chk = ttk.Checkbutton(self.scrollable_frame, text=texto, variable=self.vars[chave])
            chk.pack(anchor="w", padx=30, pady=2)
            
        ttk.Separator(self.scrollable_frame, orient='horizontal').pack(fill='x', padx=20, pady=15)

        chk_nomes = ttk.Checkbutton(self.scrollable_frame, text="Mascarar Nomes Próprios (Processamento Inteligente via NLP)", variable=self.vars["nomes"], command=self._toggle_modo_nomes)
        chk_nomes.pack(anchor="w", padx=30, pady=2)

        self.frame_modos_nome = ttk.Frame(self.scrollable_frame)
        self.frame_modos_nome.pack(anchor="w", padx=55, pady=5) 

        self.rb_iniciais = ttk.Radiobutton(self.frame_modos_nome, text="Substituir por Iniciais (A. C. S.)", variable=self.vars["modo_nomes"], value="iniciais", state="disabled")
        self.rb_iniciais.pack(anchor="w", pady=2)
        
        self.rb_total = ttk.Radiobutton(self.frame_modos_nome, text="Ocultação Total ([NOME])", variable=self.vars["modo_nomes"], value="total", state="disabled")
        self.rb_total.pack(anchor="w", pady=2)

    def confirmar(self, tipo_origem):
        opcoes = {k: v.get() for k, v in self.vars.items()}
        self.resultado = (opcoes, tipo_origem)
        self.destroy()

    def cancelar(self):
        self.resultado = None
        self.destroy()

    def _toggle_modo_nomes(self):
        if self.vars["nomes"].get():
            self.rb_iniciais.config(state="normal")
            self.rb_total.config(state="normal")
        else:
            self.rb_iniciais.config(state="disabled")
            self.rb_total.config(state="disabled")


def iniciar_interface_grafica():
    try:
        import pyi_splash
        pyi_splash.close()
    except ImportError:
        pass

    if "NUITKA_ONEFILE_PARENT" in os.environ:
        try:
            import sys
            if sys.platform == "win32":
                import ctypes
                ctypes.windll.kernel32.SetEvent(int(os.environ["NUITKA_ONEFILE_PARENT"]))
            elif sys.platform.startswith("linux") or sys.platform == "darwin":
                import signal
                os.kill(int(os.environ["NUITKA_ONEFILE_PARENT"]), signal.SIGTERM)
        except Exception as e:
            logger.warning(f"O fechamento do splash nativo não foi necessário ou falhou: {e}")

    root = tk.Tk()
    root.withdraw()

    dialogo = DialogoOpcoes(root)
    root.wait_window(dialogo)
    
    if not dialogo.resultado:
        logger.info("Operação cancelada pelo usuário na interface gráfica.")
        return

    opcoes_escolhidas, tipo_origem = dialogo.resultado
    arquivos_para_processar = []
    
    extensoes_suportadas = ('.pdf', '.txt', '.docx')

    if tipo_origem == "diretorio":
        pasta_selecionada = filedialog.askdirectory(title="Selecione o Diretório Raiz")
        if not pasta_selecionada: return

        for raiz, dirs, arqs in os.walk(pasta_selecionada):
            for arq in arqs:
                if arq.lower().endswith(extensoes_suportadas) and "_anonimizado" not in arq:
                    arquivos_para_processar.append(os.path.join(raiz, arq))
    else:
        arquivo = filedialog.askopenfilename(
            title="Selecione o Documento", 
            filetypes=[
                ("Documentos Suportados", "*.pdf;*.txt;*.docx"),
                ("PDF", "*.pdf"),
                ("Texto Puro", "*.txt"),
                ("Microsoft Word", "*.docx")
            ]
        )
        if not arquivo: return
        arquivos_para_processar.append(arquivo) 

    if not arquivos_para_processar:
        messagebox.showinfo("Info", "Nenhum arquivo suportado encontrado na seleção.")
        return

    contador = 0
    total = len(arquivos_para_processar)
    logger.info(f"Iniciando processamento. Total de arquivos: {total}")
    
    with DocumentAnonimizer(opcoes=opcoes_escolhidas) as anonimizador:
        for caminho in arquivos_para_processar:
            try:
                anonimizador.processar_arquivo(caminho)
                contador += 1
            except Exception:
                pass 
            
    mensagem_final = f"Processamento finalizado!\n{contador} de {total} arquivos gerados com sucesso."
    logger.info(mensagem_final.replace('\n', ' - '))
    messagebox.showinfo("Concluído", mensagem_final)


if __name__ == "__main__":
    iniciar_interface_grafica()